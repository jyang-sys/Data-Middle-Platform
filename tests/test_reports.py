import io
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import main, ingestion


def order(identity, day='2026-07-31T12:00:00', sku='A', quantity=2, **extra):
    return {'order_id': identity, 'date_add': int(datetime.fromisoformat(day).replace(tzinfo=timezone.utc).timestamp()),
            'date_confirmed': 1790800000, 'confirmed': True, 'order_status_id': 10,
            'order_source': 'amazon', 'order_source_id': 1,
            'products': [{'sku': sku, 'quantity': quantity}], **extra}


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db_patch = patch.object(main, 'DB_PATH', Path(self.tmp.name) / 'test.db')
        self.db_patch.start()
        self.context = TestClient(main.app)
        self.client = self.context.__enter__()
        self.params = {'start': '2026-07-01', 'end': '2026-07-31', 'timezone': 'UTC'}
        self.seed([order(1), order(2, sku='B', quantity=3, order_source='ebay'),
                   order(3, sku='C', quantity=4), order(4, quantity=9, confirmed=False),
                   order(5, quantity=8, order_source='order_return'), order(6, quantity=7),
                   order(7, quantity=6, date_add=None), order(8, day='2026-08-01T00:00:00')])
        self.annotate(6, {'exclude_from_sales': True})
        self.map('A', '128+6', '蓝色')
        self.map('B', '128+6', '蓝色', platform='ebay')
        self.map('C', '256+8', '黑色')

    def tearDown(self):
        self.context.__exit__(None, None, None)
        self.db_patch.stop()
        self.tmp.cleanup()

    def seed(self, rows):
        with main.database() as db:
            ingestion.save_order_page(db, rows, ingestion.now())

    def annotate(self, identity, extra):
        response = self.client.patch(f'/api/management/orders/{identity}/ownership', json={
            'expected_version': 0, 'actor': '验收', 'extra_fields': extra})
        self.assertEqual(response.status_code, 200, response.text)

    def map(self, sku, spec='', color='', platform='*', model='C51', version=0):
        result = self.client.put('/api/reports/mappings', json={
            'platform': platform, 'source_sku': sku, 'model': model, 'specification': spec,
            'color': color, 'product_type': '手机', 'expected_version': version, 'actor': '验收'})
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    def report(self, **changes):
        response = self.client.get('/api/reports/sales', params={**self.params, **changes})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_unit_totals_reconcile_all_dimensions_and_details(self):
        result = self.report()
        self.assertTrue(result['trial'])
        self.assertEqual((result['total_units'], result['order_count'], result['line_count']), (9, 3, 3))
        for name in ['monthly', 'daily', 'product_daily']:
            self.assertEqual(sum(r['total'] for r in result[name]), 9)
        for name in ['groups', 'employees', 'shops']:
            self.assertEqual(sum(r['quantity'] for r in result[name]), 9)
        self.assertEqual(result['groups'][0]['name'], '未分配')
        detail = self.client.get('/api/reports/details', params=self.params).json()
        self.assertEqual(detail['total_units'], 9)
        self.assertEqual(sum(r['quantity'] for r in detail['lines']), 9)
        self.assertEqual(result['monthly'][0]['total'], 5)
        self.assertEqual(len(result['daily']), 31)
        self.assertEqual(result['daily'][-1]['date'], '2026-07-31')
        self.assertEqual(result['daily'][-1]['total'], 9)
        self.assertEqual(result['diagnostics']['excluded_return_orders'], 1)
        self.assertEqual(result['diagnostics']['excluded_marked_orders'], 1)
        self.assertEqual(result['diagnostics']['missing_date_orders'], 1)

    def test_model_rollup_and_exact_filters(self):
        self.assertEqual(len(self.report(granularity='model')['monthly']), 1)
        self.assertEqual(self.report(model='C51', product_type='手机')['total_units'], 9)
        self.assertEqual(self.report(platform='ebay')['total_units'], 3)
        self.assertEqual(self.report(source_sku='C')['total_units'], 4)
        self.assertEqual(self.report(confirmed_only=False)['total_units'], 18)
        self.assertEqual(self.report(status_ids=[99])['total_units'], 0)
        self.assertEqual(self.report(status_ids=[10, 99])['total_units'], 9)

    def test_timezone_and_platform_date_do_not_fallback(self):
        self.seed([order(9, day='2026-07-31T20:00:00', quantity=5)])
        self.assertEqual(self.report()['total_units'], 14)
        self.assertEqual(self.report(timezone='Asia/Shanghai')['total_units'], 9)
        self.annotate(1, {'platform_order_date': '2026-07-31'})
        self.annotate(2, {'platform_order_date': '2026-07-31T20:00:00Z'})
        self.annotate(3, {'platform_order_date': 'invalid'})
        self.assertEqual(self.report(date_basis='platform_order')['total_units'], 5)
        self.assertEqual(self.report(date_basis='platform_order', timezone='Asia/Shanghai')['total_units'], 2)

    def test_saved_sales_ownership_used_and_live_default_does_not_reassign(self):
        group = self.client.post('/api/management/groups', json={'code': 'G', 'name': '销售组'}).json()
        emp = self.client.post('/api/management/employees', json={'code': 'E', 'name': '负责人', 'group_id': group['id']}).json()
        result = self.client.patch('/api/management/orders/1/ownership', json={'expected_version': 0,
            'actor': '验收', 'employee_id': emp['id'], 'extra_fields': {'inventory_group': '借货组'}})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(self.report(group_id=group['id'])['total_units'], 2)
        self.assertEqual(self.report(employee_id=emp['id'])['total_units'], 2)
        self.assertEqual(self.report()['groups'][0]['name'], '未分配')
        self.seed([order(1, quantity=3)])
        self.assertEqual(self.report(group_id=group['id'])['total_units'], 3)
        self.assertEqual(self.report(group_id=group['id'])['groups'][0]['name'], '销售组')

    def test_unknown_codes_do_not_merge_across_platform_and_platform_mapping_wins(self):
        self.seed([order(9, sku='UNKNOWN'), order(10, sku='UNKNOWN', order_source='ebay')])
        result = self.report()
        self.assertEqual(len(result['unmapped']), 2)
        self.map('UNKNOWN', '通用', model='公共型号')
        self.map('UNKNOWN', '专用', platform='ebay', model='专用型号')
        models = {r['model']: r['total'] for r in self.report()['monthly']}
        self.assertEqual(models['公共型号'], 2)
        self.assertEqual(models['专用型号'], 2)

    def test_mapping_conflict_audit_and_idempotence(self):
        result = self.client.put('/api/reports/mappings', json={'source_sku': 'A', 'model': '错误', 'expected_version': 0, 'actor': '验收'})
        self.assertEqual(result.status_code, 409)
        self.map('A', '128+6', '蓝色', version=1)
        history = self.client.get('/api/reports/mappings/history', params={'platform': '*', 'source_sku': 'A'}).json()
        self.assertEqual(len(history), 1)
        self.map('A', '128+6', '蓝色', model='新型号', version=1)
        history = self.client.get('/api/reports/mappings/history', params={'platform': '*', 'source_sku': 'A'}).json()
        self.assertEqual(len(history), 2)
        self.assertEqual(history[0]['before']['model'], 'C51')
        self.assertEqual(history[0]['after']['model'], '新型号')
        self.seed([order(1)])
        self.assertEqual(self.report(model='新型号')['total_units'], 2)
        raw = self.client.get('/api/base/orders/1').json()['source_order']
        self.assertEqual(raw['products'][0]['sku'], 'A')

    def test_months_keep_years_and_invalid_quantities_visible(self):
        self.seed([order(9, day='2025-07-31T12:00:00', quantity=5), order(10, sku='INVALID')])
        with main.database() as db:
            db.execute("UPDATE base_order_items SET quantity='2.5' WHERE order_id=10")
        result = self.report(start='2025-07-31')
        self.assertEqual({r['month'] for r in result['monthly']}, {'2025-07', '2026-07'})
        self.assertEqual(result['total_units'], 14)
        self.assertEqual(result['diagnostics']['invalid_quantity_lines'], 1)
        self.assertEqual(self.report(start='2026-07-29')['total_units'], 9)
        self.assertEqual(len(self.report(start='2026-07-29')['daily']), 3)

    def test_validation_and_detail_pagination(self):
        for changes in [{'start': '2026-08-01'}, {'timezone': 'INVALID'}, {'start': '2024-01-01'},
                        {'group_id': 0}, {'date_basis': 'confirmed'}, {'shop_id': '1'}]:
            self.assertEqual(self.client.get('/api/reports/sales', params={**self.params, **changes}).status_code, 422)
        result = self.report()['monthly'][0]
        scope = {**self.params, 'key_json': json.dumps(result['key']), 'month': result['month'], 'limit': 1}
        first = self.client.get('/api/reports/details', params=scope).json()
        second = self.client.get('/api/reports/details', params={**scope, 'offset': 1}).json()
        self.assertEqual(first['total_units'], 5)
        self.assertEqual(first['total'], 2)
        self.assertNotEqual(first['lines'][0]['order_id'], second['lines'][0]['order_id'])
        self.assertEqual(self.client.get('/api/reports/details', params={**scope, 'key_json': '{}'}).status_code, 422)

    def test_export_reconciles_and_canonical_keys_preserve_label_collisions(self):
        # These have the same display label but different specification/color identity.
        self.map('X', 'B', '', model='A')
        self.map('Y', '', '', model='A · B')
        self.map('Z', '', '', model='=1+1')
        self.seed([order(9, sku='X'), order(10, sku='Y'), order(11, sku='Z')])
        response = self.client.get('/api/reports/export', params=self.params)
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(wb.sheetnames), 7)
        for name in ['按平台汇总', '按型号汇总']:
            rows = list(wb[name].values)[1:]
            self.assertEqual(sum(r[3] for r in rows), 15)
            self.assertEqual(len([r for r in rows if 'A · B' in r[:3]]), 2)
        self.assertEqual(sum(r[-1] for r in list(wb['订单明细'].values)[1:]), 15)
        self.assertEqual(wb['全平台每日销量'].cell(33, 4).value, 15)
        cells = [c for sheet in wb for row in sheet for c in row if c.value == '=1+1']
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == 's' for c in cells))
