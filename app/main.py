"""A single-user, local-first data workbench."""
import csv
import io
import json
import os
import sqlite3
import time
import uuid
import zipfile
from xml.etree.ElementTree import ParseError
from contextlib import asynccontextmanager, contextmanager
from datetime import datetime, timezone
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("DATA_PLATFORM_DB", str(ROOT.parent / "data" / "platform.sqlite3")))
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100


@contextmanager
def database():
    connection = sqlite3.connect(DB_PATH, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def initialize():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with database() as connection:
        connection.execute("PRAGMA journal_mode = WAL")
        connection.execute("""CREATE TABLE IF NOT EXISTS datasets (
            id TEXT PRIMARY KEY, name TEXT NOT NULL, filename TEXT NOT NULL,
            table_name TEXT NOT NULL UNIQUE, columns_json TEXT NOT NULL,
            row_count INTEGER NOT NULL, created_at TEXT NOT NULL
        )""")
        connection.execute("""CREATE TABLE IF NOT EXISTS sales_orders (
            source TEXT NOT NULL, order_id INTEGER NOT NULL, confirmed_at INTEGER NOT NULL,
            currency TEXT NOT NULL, total TEXT NOT NULL, delivery TEXT NOT NULL,
            status_id INTEGER NOT NULL, platform TEXT NOT NULL, shop TEXT NOT NULL,
            PRIMARY KEY (source, order_id)
        )""")
        connection.execute("""CREATE TABLE IF NOT EXISTS sales_items (
            source TEXT NOT NULL, order_id INTEGER NOT NULL, position INTEGER NOT NULL,
            name TEXT NOT NULL, quantity TEXT NOT NULL, total TEXT NOT NULL, sku TEXT NOT NULL,
            PRIMARY KEY (source, order_id, position)
        )""")
        connection.execute("""CREATE TABLE IF NOT EXISTS sync_state (
            source TEXT PRIMARY KEY, updated_at TEXT NOT NULL, order_count INTEGER NOT NULL
        )""")
        connection.execute("""CREATE TABLE IF NOT EXISTS sales_sources (
            source TEXT NOT NULL, platform TEXT NOT NULL, shop TEXT NOT NULL, name TEXT NOT NULL,
            PRIMARY KEY (source, platform, shop)
        )""")
    from app.ingestion import initialize_storage
    initialize_storage()
    from app.management import initialize_storage as initialize_management
    initialize_management()
    from app.reports import initialize_storage as initialize_reports
    initialize_reports()


@asynccontextmanager
async def lifespan(app):
    initialize()
    yield


app = FastAPI(title="数据中台 · 基础版", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")


@app.middleware("http")
async def same_origin_mutations(request: Request, call_next):
    # Browsers may not submit cross-site uploads/deletions to this local tool.
    if request.method not in {"GET", "HEAD", "OPTIONS"}:
        origin = request.headers.get("origin")
        if origin and origin != str(request.base_url).rstrip("/"):
            from fastapi.responses import JSONResponse
            return JSONResponse({"detail": "请从应用页面操作。"}, status_code=403)
    return await call_next(request)


@app.get("/")
@app.get("/index.html", include_in_schema=False)
def index():
    return FileResponse(ROOT / "static" / "index.html")


@app.get("/{asset_name}", include_in_schema=False)
def root_asset(asset_name: str):
    # Relative resources work at both / and /static/index.html, as well as
    # when the HTML is opened beside its assets in a file preview.
    if asset_name not in {"style.css", "management.css", "app.js", "management.js", "reports.js", "reports.css", "inventory.js", "inventory.css"}:
        raise HTTPException(404, "资源不存在。")
    return FileResponse(ROOT / "static" / asset_name)


@app.get("/api/health")
def health():
    with database() as connection:
        connection.execute("SELECT 1").fetchone()
    return {"status": "ok"}


def dataset_dict(row):
    result = dict(row)
    result["columns"] = json.loads(result.pop("columns_json"))
    return result


def get_dataset(connection, dataset_id):
    row = connection.execute("SELECT * FROM datasets WHERE id = ?", (dataset_id,)).fetchone()
    if row is None:
        raise HTTPException(404, "数据集不存在。")
    return row


@app.get("/api/datasets")
def list_datasets():
    with database() as connection:
        return [dataset_dict(row) for row in connection.execute("SELECT * FROM datasets ORDER BY created_at DESC")]


@app.get("/api/overview")
def overview():
    with database() as connection:
        rows = connection.execute("SELECT row_count, columns_json FROM datasets").fetchall()
    return {"datasets": len(rows), "rows": sum(row["row_count"] for row in rows),
            "columns": sum(len(json.loads(row["columns_json"])) for row in rows)}


def read_table(content, filename):
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise HTTPException(400, "无法识别文件编码，请保存为 UTF-8 CSV。")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        yield from csv.reader(io.StringIO(text, newline=""), dialect, strict=True)
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise HTTPException(400, "Excel 解压后过大，请拆分后上传。")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            yield from workbook.worksheets[0].iter_rows(values_only=True)
        finally:
            workbook.close()
    else:
        raise HTTPException(400, "仅支持 .csv 和 .xlsx 文件。")


def quote_identifier(value):
    return '"' + value.replace('"', '""') + '"'


def import_table(content, filename):
    iterator = iter(read_table(content, filename))
    try:
        raw_headers = next(iterator)
    except StopIteration:
        raise HTTPException(400, "文件为空，请提供含表头的数据。")
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    if not headers or len(headers) > MAX_COLUMNS:
        raise HTTPException(400, f"表格需要 1–{MAX_COLUMNS} 列。")
    if any(not name or len(name) > 128 or "\x00" in name for name in headers):
        raise HTTPException(400, "表头不能为空、包含空字符或超过 128 个字符。")
    if len({name.casefold() for name in headers}) != len(headers):
        raise HTTPException(400, "表头不能重名（不区分大小写）。")
    dataset_id = uuid.uuid4().hex
    table_name = "data_" + dataset_id[:12]
    columns_sql = ", ".join(quote_identifier(name) + " TEXT" for name in headers)
    insert_sql = f'INSERT INTO "{table_name}" VALUES ({", ".join("?" for _ in headers)})'
    with database() as connection:
        connection.execute("BEGIN")
        connection.execute(f'CREATE TABLE "{table_name}" ({columns_sql})')
        count = 0
        batch = []
        for line, raw in enumerate(iterator, start=2):
            if not raw or all(value in (None, "") for value in raw):
                continue
            if len(raw) != len(headers):
                raise HTTPException(400, f"第 {line} 行列数与表头不一致。")
            count += 1
            if count > MAX_ROWS:
                raise HTTPException(400, f"最多支持 {MAX_ROWS:,} 行，请拆分文件。")
            batch.append([str(value) if value is not None else None for value in raw])
            if len(batch) >= 1000:
                connection.executemany(insert_sql, batch)
                batch.clear()
        if count == 0:
            raise HTTPException(400, "表格只有表头，没有数据行。")
        if batch:
            connection.executemany(insert_sql, batch)
        connection.execute("INSERT INTO datasets VALUES (?, ?, ?, ?, ?, ?, ?)",
                           (dataset_id, Path(filename).stem, filename, table_name,
                            json.dumps(headers, ensure_ascii=False), count,
                            datetime.now(timezone.utc).isoformat()))
        return dataset_dict(get_dataset(connection, dataset_id))


@app.post("/api/datasets", status_code=201)
async def upload_dataset(file: UploadFile):
    try:
        content = await file.read(MAX_BYTES + 1)
        if len(content) > MAX_BYTES:
            raise HTTPException(413, "文件不能超过 10 MB。")
        filename = Path((file.filename or "").replace("\\", "/")).name
        if not filename:
            raise HTTPException(400, "文件名不能为空。")
        # Imports run in the thread pool so parsing doesn't block the event loop.
        from starlette.concurrency import run_in_threadpool
        return await run_in_threadpool(import_table, content, filename)
    except HTTPException:
        raise
    except (csv.Error, ValueError, sqlite3.Error, zipfile.BadZipFile, KeyError, OSError,
            ParseError, InvalidFileException) as error:
        raise HTTPException(400, "文件格式错误，请检查 CSV 或 Excel 内容。") from error
    finally:
        await file.close()


@app.get("/api/datasets/{dataset_id}/preview")
def preview(dataset_id: str, limit: int = 50):
    limit = min(max(limit, 1), 200)
    with database() as connection:
        dataset = get_dataset(connection, dataset_id)
        cursor = connection.execute(f'SELECT * FROM "{dataset["table_name"]}" LIMIT ?', (limit,))
        rows = [list(row) for row in cursor.fetchall()]
    return {"columns": json.loads(dataset["columns_json"]), "rows": rows}


@app.delete("/api/datasets/{dataset_id}")
def delete_dataset(dataset_id: str):
    with database() as connection:
        dataset = get_dataset(connection, dataset_id)
        connection.execute(f'DROP TABLE "{dataset["table_name"]}"')
        connection.execute("DELETE FROM datasets WHERE id = ?", (dataset_id,))
    return {"deleted": True}


class Query(BaseModel):
    sql: str = Field(min_length=1, max_length=10_000)


@app.post("/api/query")
def query(body: Query):
    started = time.monotonic()
    # URI mode and query_only protect the database even if the authorizer changes.
    connection = sqlite3.connect(DB_PATH.resolve().as_uri() + "?mode=ro", uri=True)
    try:
        connection.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, 1_000_000)
        connection.setlimit(sqlite3.SQLITE_LIMIT_SQL_LENGTH, 10_000)
        connection.execute("PRAGMA query_only = ON")
        tables = {row[0] for row in connection.execute("SELECT table_name FROM datasets")}
        def authorize(action, arg1, arg2, db_name, trigger):
            if action == sqlite3.SQLITE_READ:
                return sqlite3.SQLITE_OK if arg1 in tables else sqlite3.SQLITE_DENY
            if action in (sqlite3.SQLITE_SELECT, sqlite3.SQLITE_RECURSIVE):
                return sqlite3.SQLITE_OK
            if action == sqlite3.SQLITE_FUNCTION and (arg2 or "").lower() != "load_extension":
                return sqlite3.SQLITE_OK
            return sqlite3.SQLITE_DENY
        connection.set_authorizer(authorize)
        connection.set_progress_handler(lambda: time.monotonic() - started > 3, 1000)
        cursor = connection.execute(body.sql)
        if cursor.description is None:
            raise HTTPException(400, "只允许查询数据集。")
        columns = [item[0] for item in cursor.description]
        rows = cursor.fetchmany(201)
        return {"columns": columns, "rows": [list(row) for row in rows[:200]],
                "truncated": len(rows) > 200, "elapsed_ms": round((time.monotonic() - started) * 1000, 1)}
    except sqlite3.Error as error:
        message = str(error)
        if "interrupted" in message:
            message = "查询超过 3 秒，请缩小范围。"
        elif "authorized" in message or "readonly" in message or "prohibited" in message:
            message = "只允许对已上传的数据集执行只读查询。"
        raise HTTPException(400, message) from error
    finally:
        connection.close()


from app.sales import router as sales_router
app.include_router(sales_router)
from app.ingestion import router as ingestion_router
app.include_router(ingestion_router)
from app.management import router as management_router
app.include_router(management_router)

from app.reports import router as reports_router
app.include_router(reports_router)
from app.stock_dashboard import router as stock_dashboard_router
app.include_router(stock_dashboard_router)
