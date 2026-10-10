"""Trial unit-sales reports derived from one shared order-line selection."""
import io
import json
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from typing import Annotated
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError, model_validator

from app.ingestion import encode, now

router = APIRouter(prefix='/api/reports', tags=['销量试算报表'])
SCHEMA = """
CREATE TABLE IF NOT EXISTS sales_sku_mappings (
    platform TEXT NOT NULL, source_sku TEXT NOT NULL, model TEXT NOT NULL,
    specification TEXT NOT NULL, color TEXT NOT NULL, product_type TEXT NOT NULL,
    version INTEGER NOT NULL, updated_at TEXT NOT NULL,
    PRIMARY KEY(platform,source_sku)
);
CREATE TABLE IF NOT EXISTS sales_mapping_history (
    id INTEGER PRIMARY KEY, platform TEXT NOT NULL, source_sku TEXT NOT NULL,
    before_json TEXT NOT NULL, after_json TEXT NOT NULL, actor TEXT NOT NULL,
    reason TEXT NOT NULL, created_at TEXT NOT NULL
);
"""
PLATFORM_NAMES = {'amazon': '亚马逊', 'ebay': 'eBay', 'tiktok': 'TikTok', 'shop': '独立站',
                  'walmart': '沃尔玛', 'allegro': 'Allegro', 'temu': 'TEMU', 'shein': 'SHEIN',
                  'cdiscount': 'CD', 'worten': 'WORTEN', 'debenhams': 'Debenhams'}


def initialize_storage():
    from app.main import database
    with database() as db:
        db.executescript(SCHEMA)


def platform_name(code):
    return PLATFORM_NAMES.get(code, code or '未提供平台')


class MappingInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    platform: str = Field(default='*', min_length=1, max_length=100)
    source_sku: str = Field(min_length=1, max_length=200)
    model: str = Field(min_length=1, max_length=100)
    specification: str = Field(default='', max_length=100)
    color: str = Field(default='', max_length=100)
    product_type: str = Field(default='未分类', min_length=1, max_length=100)
    expected_version: int = Field(ge=0)
    actor: str = Field(min_length=1, max_length=100)
    reason: str = Field(default='', max_length=500)

    @model_validator(mode='after')
    def trim_fields(self):
        for name in ('platform', 'source_sku', 'model', 'specification', 'color', 'product_type', 'actor', 'reason'):
            setattr(self, name, getattr(self, name).strip())
        if any(not getattr(self, name) for name in ('platform', 'source_sku', 'model', 'product_type', 'actor')):
            raise ValueError('型号、SKU、产品类型、平台和操作人不能为空')
        return self


@router.get('/mappings')
def mappings(limit: int = Query(100, ge=1, le=200), offset: int = Query(0, ge=0), search: str = ''):
    from app.main import database
    with database() as db:
        # instr treats %/_ literally; product codes are not SQL patterns.
        condition = ' WHERE instr(source_sku,?)>0 OR instr(model,?)>0'
        total = db.execute('SELECT COUNT(*) FROM sales_sku_mappings' + condition, (search, search)).fetchone()[0]
        rows = [dict(row) for row in db.execute('SELECT * FROM sales_sku_mappings' + condition +
                ' ORDER BY platform,source_sku LIMIT ? OFFSET ?', (search, search, limit, offset))]
    return {'total': total, 'mappings': rows}


@router.put('/mappings')
def save_mapping(body: MappingInput):
    from app.main import database
    with database() as db:
        db.execute('BEGIN IMMEDIATE')
        old = db.execute('SELECT * FROM sales_sku_mappings WHERE platform=? AND source_sku=?',
                         (body.platform, body.source_sku)).fetchone()
        version = old['version'] if old else 0
        if version != body.expected_version:
            raise HTTPException(409, '映射已被修改，请重新读取再保存。')
        values = body.model_dump(exclude={'expected_version', 'actor', 'reason'})
        if old and all(old[key] == value for key, value in values.items()):
            return dict(old)
        db.execute('''INSERT INTO sales_sku_mappings VALUES(?,?,?,?,?,?,?,?)
            ON CONFLICT(platform,source_sku) DO UPDATE SET model=excluded.model,
            specification=excluded.specification,color=excluded.color,product_type=excluded.product_type,
            version=excluded.version,updated_at=excluded.updated_at''',
                   (body.platform, body.source_sku, body.model, body.specification, body.color,
                    body.product_type, version + 1, now()))
        saved = dict(db.execute('SELECT * FROM sales_sku_mappings WHERE platform=? AND source_sku=?',
                               (body.platform, body.source_sku)).fetchone())
        db.execute('''INSERT INTO sales_mapping_history(platform,source_sku,before_json,after_json,actor,reason,created_at)
            VALUES(?,?,?,?,?,?,?)''', (body.platform, body.source_sku, encode(dict(old) if old else {}),
                                     encode(saved), body.actor, body.reason, now()))
    return saved


@router.get('/mappings/history')
def mapping_history(platform: str, source_sku: str):
    from app.main import database
    with database() as db:
        rows = db.execute('''SELECT * FROM sales_mapping_history WHERE platform=? AND source_sku=?
            ORDER BY id DESC LIMIT 100''', (platform, source_sku))
        result = []
        for row in rows:
            value = dict(row)
            value['before'] = json.loads(value.pop('before_json'))
            value['after'] = json.loads(value.pop('after_json'))
            result.append(value)
    return result


class ReportFilters(BaseModel):
    start: date | None = None
    end: date | None = None
    timezone: str = 'Asia/Shanghai'
    date_basis: str = Field(default='base_created', pattern='^(base_created|platform_order)$')
    confirmed_only: bool = True
    status_ids: list[int] = Field(default_factory=list)
    platform: str | None = None
    shop_id: str | None = None
    group_id: int | None = Field(default=None, ge=1)
    employee_id: int | None = Field(default=None, ge=1)
    model: str | None = None
    source_sku: str | None = None
    product_type: str | None = None
    granularity: str = Field(default='specification', pattern='^(model|specification|source_product)$')

    @model_validator(mode='after')
    def validate_range(self):
        try:
            zone = ZoneInfo(self.timezone)
        except (ZoneInfoNotFoundError, ValueError) as error:
            raise ValueError('请选择有效的 IANA 时区') from error
        self.end = self.end or datetime.now(zone).date()
        self.start = self.start or self.end - timedelta(days=6)
        if self.end < self.start or (self.end - self.start).days > 365:
            raise ValueError('开始日期不能晚于结束日期，范围最多为 366 天')
        if self.shop_id is not None and self.platform is None:
            raise ValueError('筛选店铺时必须同时选择平台')
        return self


def report_filters(start: date | None = None, end: date | None = None,
                   timezone: str = 'Asia/Shanghai', date_basis: str = 'base_created',
                   confirmed_only: bool = True, status_ids: list[int] = Query(default=[]),
                   platform: str | None = None, shop_id: str | None = None,
                   group_id: int | None = None, employee_id: int | None = None,
                   model: str | None = None, source_sku: str | None = None,
                   product_type: str | None = None, granularity: str = 'specification'):
    try:
        return ReportFilters.model_validate(locals())
    except ValidationError as error:
        raise HTTPException(422, '; '.join(e['msg'] for e in error.errors())) from error


def selected_day(row, filters, zone):
    if filters.date_basis == 'base_created':
        value = row['created_at']
        if value is None or value <= 0:
            return None
        try:
            return datetime.fromtimestamp(value, timezone.utc).astimezone(zone).date()
        except (OverflowError, ValueError, OSError):
            return None
    value = json.loads(row['extra_json'] or '{}').get('platform_order_date')
    if not isinstance(value, str):
        return None
    try:
        timestamp = datetime.fromisoformat(value.replace('Z', '+00:00'))
        # A date or local timestamp explicitly uses the selected report timezone.
        if timestamp.tzinfo is None:
            timestamp = timestamp.replace(tzinfo=zone)
        return timestamp.astimezone(zone).date()
    except ValueError:
        return None


def report_lines(filters):
    from app.main import database
    zone = ZoneInfo(filters.timezone)
    with database() as db:
        db.execute('BEGIN')  # All dimensions derive from one source snapshot.
        mapping_rows = db.execute('SELECT * FROM sales_sku_mappings').fetchall()
        mappings = {(r['platform'], r['source_sku']): dict(r) for r in mapping_rows}
        orders = db.execute('''SELECT o.order_id,o.external_order_id,o.platform,o.shop_id,o.created_at,
            o.confirmed,o.status_id,o.updated_at,a.employee_id,a.group_id,a.sales_owner,a.business_group,
            a.extra_json,s.name AS shop_name FROM base_orders o LEFT JOIN order_annotations a
            ON a.order_id=o.order_id LEFT JOIN base_shops s ON s.platform=o.platform AND s.shop_id=o.shop_id''').fetchall()
        items = defaultdict(list)
        for row in db.execute('SELECT order_id,position,sku,quantity,storage,storage_id,product_id,variant_id,raw_json FROM base_order_items'):
            items[row['order_id']].append(dict(row))
        products = {(str(r['inventory_id']), str(r['product_id'])): dict(r) for r in db.execute('SELECT inventory_id,product_id,parent_id,sku,name FROM base_products')}
        sync = db.execute("SELECT finished_at FROM base_sync_jobs WHERE domain='orders' AND status='success' ORDER BY started_at DESC LIMIT 1").fetchone()
    diagnostics = {'missing_date_orders': 0, 'excluded_return_orders': 0, 'excluded_marked_orders': 0,
                   'invalid_quantity_lines': 0, 'unmapped_lines': 0, 'orders_without_items': 0}
    lines = []
    for order in orders:
        if filters.platform is not None and order['platform'] != filters.platform:
            continue
        if filters.shop_id is not None and order['shop_id'] != filters.shop_id:
            continue
        if filters.group_id is not None and order['group_id'] != filters.group_id:
            continue
        if filters.employee_id is not None and order['employee_id'] != filters.employee_id:
            continue
        if filters.confirmed_only and not order['confirmed']:
            continue
        if filters.status_ids and order['status_id'] not in filters.status_ids:
            continue
        day = selected_day(order, filters, zone)
        if day is None:
            diagnostics['missing_date_orders'] += 1
            continue
        if not filters.start <= day <= filters.end:
            continue
        if order['platform'] == 'order_return':
            diagnostics['excluded_return_orders'] += 1
            continue
        extra = json.loads(order['extra_json'] or '{}')
        if extra.get('exclude_from_sales') is True:
            diagnostics['excluded_marked_orders'] += 1
            continue
        if not items[order['order_id']]:
            diagnostics['orders_without_items'] += 1
        for item in items[order['order_id']]:
            if filters.source_sku is not None and item['sku'] != filters.source_sku:
                continue
            try:
                quantity = Decimal(item['quantity'])
                if not quantity.is_finite() or quantity < 0 or quantity != quantity.to_integral_value():
                    raise ValueError('invalid unit count')
            except (InvalidOperation, ValueError, TypeError):
                diagnostics['invalid_quantity_lines'] += 1
                continue
            product_id = item['variant_id'] if item['variant_id'] not in (None, '', '0') else item['product_id']
            product = products.get((item['storage_id'], product_id)) if item['storage'] == 'db' else None
            if product and product_id != item['product_id'] and str(product['parent_id']) != item['product_id']:
                product = None
            source_name = str((product['name'] if product and product['name'] not in ('', 'None') else json.loads(item['raw_json']).get('name', '')) or item['sku'] or '未提供名称')
            parent = products.get((item['storage_id'], str(product['parent_id']))) if product and product['parent_id'] else product
            mapping = mappings.get((order['platform'], item['sku'])) or mappings.get(('*', item['sku']))
            if mapping:
                model = mapping['model']
                specification = mapping['specification']
                color = mapping['color']
                product_type = mapping['product_type']
                label = model if filters.granularity == 'model' else ' · '.join(x for x in (model, specification, color) if x)
            else:
                model = '未映射 SKU: ' + (item['sku'] or '未提供 SKU') + ' [' + order['platform'] + ']'
                specification = color = ''
                product_type = '未分类'
                label = model
            if filters.granularity == 'source_product':
                label = source_name + ' · ' + (item['sku'] or '无 SKU')
            if filters.model is not None and model != filters.model:
                continue
            if filters.product_type is not None and product_type != filters.product_type:
                continue
            if mapping is None:
                diagnostics['unmapped_lines'] += 1
            # Tuple keys prevent collisions between different configurations with the same label.
            key = (model,) if filters.granularity == 'model' else (model, specification, color)
            if filters.granularity == 'source_product':
                key = ('source', str(product['inventory_id']), str(product['product_id'])) if product else ('unlinked', order['platform'], item['sku'] or f"{order['order_id']}:{item['position']}")
            lines.append({'order_id': order['order_id'], 'external_order_id': order['external_order_id'],
                'position': item['position'], 'date': day.isoformat(), 'month': day.strftime('%Y-%m'),
                'platform': order['platform'], 'platform_name': platform_name(order['platform']),
                'shop_id': order['shop_id'], 'shop_name': order['shop_name'] or order['shop_id'],
                'employee_id': order['employee_id'], 'group_id': order['group_id'],
                'sales_owner': order['sales_owner'], 'business_group': order['business_group'],
                'source_sku': item['sku'], 'source_name': source_name, 'parent_sku': parent['sku'] if parent else None,
                'inventory_id': product['inventory_id'] if product else None, 'product_id': product['product_id'] if product else None, 'model': model, 'specification': specification, 'color': color,
                'product_type': product_type, 'label': label, 'quantity': int(quantity), 'mapped': mapping is not None,
                'key': key})
    return lines, diagnostics, sync['finished_at'] if sync else None


def aggregate(filters):
    lines, diagnostics, last_sync = report_lines(filters)
    platforms = sorted({r['platform'] for r in lines})
    monthly = {}
    daily = { (filters.start + timedelta(days=i)).isoformat(): defaultdict(int)
              for i in range((filters.end - filters.start).days + 1) }
    product_daily = {}
    breakdowns = {name: {} for name in ('groups', 'employees', 'shops')}
    type_sets = defaultdict(set)
    for row in lines:
        q = row['quantity']
        type_sets[row['key']].add(row['product_type'])
        key = (row['month'], row['key'])
        target = monthly.setdefault(key, {'month': row['month'], 'model': row['model'],
            'label': row['label'], 'key': list(row['key']), 'platforms': defaultdict(int), 'total': 0})
        target['platforms'][row['platform']] += q
        target['total'] += q
        daily[row['date']][row['platform']] += q
        product_key = (row['platform'], row['key'])
        target = product_daily.setdefault(product_key, {'platform': row['platform'], 'label': row['label'],
            'model': row['model'], 'key': list(row['key']), 'daily': defaultdict(int), 'total': 0})
        target['daily'][row['date']] += q
        target['total'] += q
        for dimension, identity, label in [
            ('groups', ('id', row['group_id']) if row['group_id'] is not None else ('text', row['business_group']), row['business_group']),
            ('employees', ('id', row['employee_id']) if row['employee_id'] is not None else ('text', row['sales_owner']), row['sales_owner']),
            ('shops', (row['platform'], row['shop_id']), row['shop_name'])]:
            # IDs unify employees across name snapshots; legacy names remain separate.
            group = breakdowns[dimension].setdefault(identity, {'id': identity[1],
                'identity_type': identity[0], 'name': label or '未分配', 'quantity': 0})
            group['quantity'] += q
    for row in list(monthly.values()) + list(product_daily.values()):
        row['product_type'] = ' / '.join(sorted(type_sets[tuple(row['key'])]))
    daily_rows = [{'date': day, 'platforms': dict(values), 'total': sum(values.values())} for day, values in daily.items()]
    warnings = ['试算口径：取消、退款及退货扣减规则尚未确认；结果不作为最终净销量。',
                'Base 来源 order_return 和本地附加字段 exclude_from_sales=true 的订单已排除。',
                '销售小组与负责人采用已保存的订单归属，不以库存来源或当前店铺默认值替代。',
                '当前映射用于重新计算历史报表；映射修改会影响历史查询，尚无结账冻结。']
    warnings.append('日期采用 Base 创建时间，尚未认定为平台原始下单时间。' if filters.date_basis == 'base_created'
                    else '日期采用订单本地附加字段 platform_order_date，未填写或格式无效的订单不计入。')
    if diagnostics['unmapped_lines']:
        warnings.append('业务规格尚未映射；系统产品按明确的目录/产品 ID 汇总，不认定为统一业务规格。' if filters.granularity == 'source_product' else '未映射商品保留原 SKU 单独汇总，不猜测型号或拆分合并配置。')
    unmapped = {}
    for row in lines:
        if not row['mapped']:
            target = unmapped.setdefault((row['platform'], row['source_sku']), {'platform': row['platform'], 'source_sku': row['source_sku'], 'quantity': 0})
            target['quantity'] += row['quantity']
    return {'trial': True, 'unmapped': sorted(unmapped.values(), key=lambda r: -r['quantity']), 'filters': filters.model_dump(mode='json'), 'unit': '件',
            'total_units': sum(r['quantity'] for r in lines), 'order_count': len({r['order_id'] for r in lines}),
            'line_count': len(lines), 'platforms': [{'id': p, 'name': platform_name(p)} for p in platforms],
            'monthly': sorted(monthly.values(), key=lambda r: (r['month'], -r['total'], r['label'])),
            'daily': daily_rows, 'product_daily': sorted(product_daily.values(), key=lambda r: (r['platform'], -r['total'], r['label'])),
            **{name: sorted(values.values(), key=lambda r: (-r['quantity'], r['name'])) for name, values in breakdowns.items()},
            'diagnostics': diagnostics, 'warnings': warnings, 'last_sync': last_sync}, lines


@router.get('/sales')
def sales(filters: Annotated[ReportFilters, Depends(report_filters)]):
    return aggregate(filters)[0]


@router.get('/details')
def details(filters: Annotated[ReportFilters, Depends(report_filters)], limit: int = Query(50, ge=1, le=200),
            offset: int = Query(0, ge=0), month: str | None = None, label: str | None = None,
            key_json: str | None = None):
    lines, _, _ = report_lines(filters)
    key = None
    if key_json is not None:
        try:
            value = json.loads(key_json)
            if not isinstance(value, list) or len(value) not in (1, 3) or not all(isinstance(x, str) for x in value):
                raise ValueError('invalid key')
            key = tuple(value)
        except (ValueError, TypeError):
            raise HTTPException(422, '商品维度键格式不正确。')
    lines = [row for row in lines if (month is None or row['month'] == month)
             and (label is None or row['label'] == label) and (key is None or row['key'] == key)]
    lines.sort(key=lambda r: (r['date'], r['order_id'], r['position']), reverse=True)
    return {'total': len(lines), 'total_units': sum(r['quantity'] for r in lines),
            'lines': [{k: v for k, v in row.items() if k != 'key'} for row in lines[offset:offset + limit]]}


@router.get('/options')
def options():
    from app.main import database
    with database() as db:
        platforms = [row[0] for row in db.execute('SELECT DISTINCT platform FROM base_orders ORDER BY platform')]
        statuses = [row[0] for row in db.execute('SELECT DISTINCT status_id FROM base_orders WHERE status_id IS NOT NULL ORDER BY status_id')]
        products = [dict(r) for r in db.execute('''SELECT DISTINCT model,product_type FROM sales_sku_mappings ORDER BY model''')]
        quality = dict(db.execute("SELECT COUNT(*) AS tagged_test_products FROM base_products WHERE UPPER(sku) LIKE 'DEMO-%' OR UPPER(sku) LIKE 'SHOWCASE-%' OR UPPER(sku) LIKE 'TEST-%'").fetchone())
        bounds = dict(db.execute('SELECT MIN(created_at) AS first_created,MAX(created_at) AS last_created,MAX(CASE WHEN confirmed=1 THEN created_at END) AS last_confirmed_created FROM base_orders WHERE created_at>0').fetchone())
        shops = [dict(r) for r in db.execute('SELECT platform,shop_id,name FROM base_shops ORDER BY platform,shop_id')]
    return {'platforms': [{'id': p, 'name': platform_name(p)} for p in platforms], 'statuses': statuses,
            'products': products, 'shops': shops, 'order_dates': bounds, 'source_quality': quality}


def append_row(sheet, values):
    sheet.append(values)
    # Product codes and names may start with =; keep them text in exported Excel.
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


@router.get('/export')
def export(filters: Annotated[ReportFilters, Depends(report_filters)]):
    report, lines = aggregate(filters)
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers, columns in [
        ('按平台汇总', ['月份', '平台', '型号 / 规格', '销量', '型号', '配置', '颜色'], ('month', 'platform_name', 'label')),
        ('按型号汇总', ['型号 / 规格', '月份', '平台', '销量', '型号', '配置', '颜色'], ('label', 'month', 'platform_name'))]:
        sheet = wb.create_sheet(name)
        append_row(sheet, headers)
        values = {}
        for row in lines:
            key = (row['month'], row['platform'], row['key'])
            target = values.setdefault(key, {'row': row, 'quantity': 0})
            target['quantity'] += row['quantity']
        for key, target in sorted(values.items()):
            row = target['row']
            spec, color = ('', '') if filters.granularity == 'model' else (row['specification'], row['color'])
            append_row(sheet, [*[row[c] for c in columns], target['quantity'], row['model'], spec, color])
    platforms = report['platforms']
    sheet = wb.create_sheet('型号平台销量')
    append_row(sheet, ['月份', '产品类型', '型号 / 规格', *[p['name'] for p in platforms], '总计'])
    for row in report['monthly']:
        append_row(sheet, [row['month'], row['product_type'], row['label'], *[row['platforms'].get(p['id'], 0) for p in platforms], row['total']])
    sheet = wb.create_sheet('按平台统计')
    days = [r['date'] for r in report['daily']]
    append_row(sheet, ['平台', '产品类型', '型号 / 规格', '总销量', *days])
    for row in report['product_daily']:
        append_row(sheet, [platform_name(row['platform']), row['product_type'], row['label'], row['total'], *[row['daily'].get(day, 0) for day in days]])
    sheet = wb.create_sheet('全平台每日销量')
    append_row(sheet, ['日期', *[p['name'] for p in platforms], '合计'])
    for row in report['daily']:
        append_row(sheet, [row['date'], *[row['platforms'].get(p['id'], 0) for p in platforms], row['total']])
    append_row(sheet, ['期间合计', *[sum(r['platforms'].get(p['id'], 0) for r in report['daily']) for p in platforms], report['total_units']])
    sheet = wb.create_sheet('订单明细')
    fields = ['date', 'order_id', 'external_order_id', 'platform_name', 'shop_name', 'sales_owner',
              'business_group', 'source_sku', 'model', 'specification', 'color', 'product_type', 'quantity']
    append_row(sheet, ['订单日期', 'Base 订单 ID', '平台订单号', '平台', '店铺', '销售负责人', '业务小组',
                       '原 SKU', '型号', '配置', '颜色', '产品类型', '销量', '来源产品名称', '父 SKU', '目录 ID', '产品 ID'])
    for row in sorted(lines, key=lambda r: (r['date'], r['order_id'], r['position'])):
        append_row(sheet, [*[row[field] for field in fields], row['source_name'], row['parent_sku'], row['inventory_id'], row['product_id']])
    sheet = wb.create_sheet('口径说明')
    append_row(sheet, ['项目', '取值'])
    for key, value in report['filters'].items():
        append_row(sheet, [key, encode(value)])
    for warning in report['warnings']:
        append_row(sheet, ['试算说明', warning])
    for key, value in report['diagnostics'].items():
        append_row(sheet, [key, value])
    for sheet in wb:
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(36, max(12, max(len(str(c.value or '')) for c in column[:100]) + 2))
    stream = io.BytesIO()
    wb.save(stream)
    return Response(stream.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': 'attachment; filename="sales-trial-report.xlsx"'})
